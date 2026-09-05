# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r'C:\Users\Yanzongning\Desktop\稿子\资助政策报告'
PHOTO = os.path.join(BASE, 'zphotos')

def set_run_font(run, name='宋体', size=11, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)

def clear_cell(cell):
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)

def set_cell(cell, text, name='宋体', size=11, bold=False, align=None):
    clear_cell(cell)
    for line in text.split('\n'):
        p = cell.add_paragraph()
        if align is not None:
            p.alignment = align
        if line == '':
            continue
        run = p.add_run(line)
        set_run_font(run, name, size, bold)

def find_row(table, keyword):
    kw = ''.join(keyword.split())
    for i, row in enumerate(table.rows):
        for c in row.cells:
            if kw in ''.join(c.text.split()):
                return i
    return None

def content_cell(table, keyword):
    """返回该标签行中用于填写内容（合并）的单元格，通常是第2格"""
    i = find_row(table, keyword)
    if i is None:
        raise ValueError('未找到行: ' + keyword)
    return table.rows[i].cells[1]

# =================== 权威团队信息（取自附件1 团队申报表） ===================
team_name = '山东药品食品职业学院“资润启梦”志愿服务队'
leader = '姚珂，18462361470'                      # 负责人（附件1）
advisor = '董怡静（专职辅导员、主管系部资助工作），18853859982'  # 指导教师（附件1）
vol_size = '10'                                  # 志愿者人数（附件1）
priority_count = '7'                             # 符合优先条件人数（附件1：7人获国/省政府奖学金或助学金）
members = '姚珂（负责人）、纪文慧、杨雨萌、郭思彤等10人'

# =================== 附件3 内容 ===================
basic = ('山东药品食品职业学院“资润启梦”志愿服务队成立于2026年6月，隶属质量管理系，由10名青年志愿者组成，'
         '其中7人符合优先条件（曾获得国家助学金或国家励志奖学金），4人具备新媒体运营、摄影摄像、文案写作等特长。'
         '2026年7月12日至19日，服务队紧扣高考志愿填报关键节点，奔赴菏泽市单县、曹县及菏泽城区，'
         '依托当地三场高考招生咨询会开展“学生资助政策进万家”宣讲。队伍秉持“让好政策从纸面走进心里”的初心，'
         '将政策宣讲与数字化服务相结合，面向高三毕业生及家长精准解读生源地助学贷款、国家助学金、国家励志奖学金等政策，累计服务1500余人。')

train = ('出发前，服务队开展专题培训：系统研读《山东省学生资助资金管理办法》等政策文件，统一宣讲口径与话术，'
         '模拟家长常见疑问并演练应答；设计153份线上问卷用于认知摸排。实行“负责人统筹+分场负责”管理：'
         '负责人姚珂负总责并统筹三场活动，纪文慧、杨雨萌、郭思彤分片包干单县、曹县、菏泽会展中心现场宣讲，'
         '团队技术骨干自主开发AI政策宣讲小程序并负责宣传文稿。每日活动结束召开复盘会，梳理问题与典型案例；'
         '严格执行安全报备与考勤制度，确保志愿服务规范、有序、可追溯。')

main = (
'一、摸排先行，扫清政策认知盲区\n'
'出发前，服务队完成153份问卷调研，结果显示不少高三毕业生和家长对资助政策“听说过、但不清楚”，'
'甚至将“无偿资助”与“有偿贷款”混为一谈而不敢申请。找准盲区才能对症下药：队员结合受助经历，把条款变成故事，'
'“不讲大道理，就讲我自己”，把从学前到高等教育的全学段资助体系梳理成通俗“大白话”，为现场宣讲打好底子。\n'
'二、现身说法，把政策讲进群众心里\n'
'7月12日单县一中、7月13日曹县三桐中学、7月19日菏泽会展中心，三场高招会现场，同样的政策对话一遍遍上演。'
'面对“助学金和助学贷款到底啥不一样”的核心疑问，队员姚珂打了个比方：“助学金是国家补贴生活费、不用还；'
'助学贷款是向国家借的学费预付款、毕业再还——一个帮现在，一个管长远。”队员杨雨萌则把申请国家励志奖学金的过程'
'拆解成一步步可操作的指引。据7月28日教育部新闻发布会，我国助学贷款已形成以生源地贷款为主体、校园地贷款为补充的模式，'
'确保“应助尽助”。讲哑的嗓子、湿透的队服，换来家长一句句“心里有底了”。\n'
'三、三大创新，打通政策落地“最后一公里”\n'
'一是把高考招生咨询会变成资助宣传主阵地，并以AI小程序作为现场赋能抓手。团队紧扣志愿填报关键节点，'
'在单县一中、曹县三桐中学、菏泽会展中心三场高招会现场设点咨询、一对一答疑、案例讲解，'
'由曾获国家资助的志愿者现身说法，重点解读生源地助学贷款、新生入学资助、困难认定及奖助学金等政策；'
'现场同步摆放“资润启梦”AI政策小程序二维码，考生和家长扫码即可即时获取政策指引与常见问题解答，'
'把线下咨询与线上智能服务无缝衔接。服务前置到入学准备阶段，变“被动等待咨询”为“主动靠前服务”，切实化解学费与就学顾虑。\n'
'二是自主开发AI政策宣讲小程序，以数字化手段延展服务时空。小程序内置资助政策宣讲内容与常见问题自动回复，'
'考生和家长可随时随地获取全天候、即时性线上指引，与线下现场服务形成互补，进一步扩大政策传播覆盖面与便捷度，'
'让“政策找人”突破场次与地域限制。\n'
'三是集中服务期结束后延伸服务链条，团队成员申请进入当地资助中心，参与暑期助学贷款勤工助学实践，'
'协助办理贷款受理、续贷及档案整理等日常业务，在一线窗口精准服务有需求的家庭。三大创新层层递进，'
'线上线下结合、集中服务与长效实践并重。\n'
'经验与长效机制\n'
'实践证明，“问卷摸排—通俗宣讲—数字赋能—一线实践”的闭环路径切实有效。队伍已形成可复制的宣讲话术库与'
'数字化工具（AI小程序），并与高中、资助中心建立常态化联系，为后续“资助政策进万家”持续行动奠定基础。'
)

problems = ('存在问题：一是基层家长政策知晓率仍偏低，部分家庭对“应助尽助”缺乏信心；'
            '二是方言与文化程度差异造成理解门槛，单次宣讲难以全覆盖；三是集中服务后缺乏长期跟踪，政策落地效果难持续。'
            '意见建议：建立“一村一档”困难学生台账，实现精准回访；扩大AI小程序覆盖面并开通图文版指引；'
            '推动与高中、乡镇资助站建立常态化宣讲机制，把集中服务延伸为全年陪伴。')

cases = (
'案例一：一场高招会上的“政策区别课”——把条文讲成听得懂的家常\n'
'7月12日，单县一中高招会现场，闷热的午后挤满了揣着分数、选着前程的考生和家长。一位皮肤黝黑的父亲替女儿问出了最忐忑的问题：'
'“这学上下来得花多少钱？家里能不能扛得住？”他的顾虑，也是山东药品食品职业学院“资润启梦”志愿服务队此行菏泽最想打消的。'
'队伍由纪文慧、姚珂、杨雨萌、郭思彤等队员组成，带着一个明确使命：让国家的好政策，从纸面上走进人心里。\n'
'准大一新生小刘紧接着追问出最核心的困惑：“学姐，助学金和助学贷款到底啥不一样？我该申请哪个？”'
'队员姚珂没有照本宣科地背文件，而是打了个朴素又准确的比方：“助学金是国家和学校补贴你的生活费，不用还，让你安心吃饭、买书；'
'助学贷款是你向国家借的学费预付款，毕业工作后再分期还。一个是帮现在，一个是管长远。”话音未落，一旁替女儿来的父亲松开了紧皱的眉头。'
'这堂“政策区别课”精准命中了群众最关心的“钱从哪来、要不要还”两个核心，把抽象条文变成听得懂、用得上的生活语言。'
'队员纪文慧更结合自己拿过国家助学金的经历现身说法：“不讲大道理，就讲我自己，大家听了就懂。”'
'据7月28日教育部新闻发布会，我国助学贷款已形成以生源地贷款为主体、校园地贷款为补充的发展模式，确保“应助尽助”'
'——这意味着无论考生户籍在哪、考到哪里，都有对应的贷款渠道兜底，学费不再是迈不进校门的坎。\n'
'案例二：从“听不懂”到“心里有底”——数字赋能让服务突破场次\n'
'在随后的曹县三桐中学高招会和菏泽会展中心，这样的对话一遍遍上演。队员杨雨萌把自己申请国家励志奖学金的全过程'
'拆解成一步步可操作的指引，从表格怎么填到找谁盖章，讲得明明白白；一位母亲听完，拍了拍女儿的肩膀：'
'“听姐姐这么说，咱就放心了。你只管好好学，别的有国家呢。”\n'
'本次服务的最大创新，是团队自主开发的AI政策宣讲小程序。考生和家长现场扫码，即可获取全天候、即时性的政策指引与常见问题自动回复，'
'把“面对面”延展为“随时随地的陪伴”，让“政策找人”突破场次与地域限制。夕阳西下，最后一场咨询会收摊，'
'一个听过讲解的女生跑回来，往队员手里塞了两瓶水，说了句“姐姐们辛苦了”便转身跑开。两瓶水，是群众最朴素的认可。\n'
'三场活动横跨单县、曹县和菏泽城区，队员顶着烈日发出数百份折页、面对面聊了上千人。'
'让每一分好政策变成普通人家看得见、摸得着的希望——这正是“青春筑梦”最实在的注脚：'
'用面对面答疑消解顾虑，用数字化工具延伸服务，用亲身经历印证政策温度，彻彻底底打通政策落地的“最后一公里”。'
)

photos = [
    (os.path.join(PHOTO, 'photo_1.jpeg'), '图1：“资润启梦”志愿服务队在单县一中高招会现场'),
    (os.path.join(PHOTO, 'photo_2.jpeg'), '图2：队员郭思彤（右）在曹县三桐中学高招会现场宣讲资助政策'),
    (os.path.join(PHOTO, 'photo_3.jpeg'), '图3：队员杨雨萌（右）在曹县高招会现场为家长讲解政策'),
    (os.path.join(PHOTO, 'photo_4.jpeg'), '图4：服务队在菏泽市招生考试服务中心前合影'),
]

print('== 处理附件3 ==')
doc3 = Document(os.path.join(BASE, '关于开展2026年山东省“学生资助政策进万家_9-10.docx'))
t0 = doc3.tables[0]
t1 = doc3.tables[1]

set_cell(content_cell(t0, '团队名称'), team_name)
set_cell(content_cell(t0, '负责人姓名及联'), leader)
set_cell(content_cell(t0, '团队基本情况'), basic)
# 开展活动场次行：两对 标签/值
ri = find_row(t0, '开展活动场次')
r = t0.rows[ri].cells
set_cell(r[1], '3')
set_cell(r[3], '1500余人')
set_cell(content_cell(t0, '集中服务时间'), '2026年7月12日—7月19日')
set_cell(content_cell(t0, '入户数'), '0（本次以高招会集中宣讲为主，未开展入户走访）')
set_cell(content_cell(t0, '相关宣传报道'), '［待补充：校院两级官网/公众号等宣传报道链接］')
set_cell(content_cell(t0, '团队开展培训及日常管理情况'), train)
set_cell(content_cell(t0, '主要服务内容及成效经验'), main)

set_cell(content_cell(t1, '存在问题及意见建议'), problems)
set_cell(content_cell(t1, '典型案例'), cases)

# 活动照片
pcell = content_cell(t1, '活动照片')
clear_cell(pcell)
for path, cap in photos:
    p = pcell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Cm(13))
    cp = pcell.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cp.add_run(cap)
    set_run_font(cr, '黑体', 10.5)

out3 = os.path.join(BASE, '附件3_优秀团队总结表_已填.docx')
doc3.save(out3)
print('保存', out3)

# =================== 附件2 内容 ===================
print('== 处理附件2 ==')
doc2 = Document(os.path.join(BASE, '关于开展2026年山东省“学生资助政策进万家_8.docx'))
t = doc2.tables[0]
# 找数据行（首格为 '1'）
row1 = None
for row in t.rows:
    if row.cells[0].text.strip() == '1':
        row1 = row
        break
if row1 is None:
    raise ValueError('未找到数据行')
c = row1.cells
set_cell(c[1], team_name)
set_cell(c[2], vol_size)                       # 团队成员人数 = 10
set_cell(c[3], members)                         # 团队成员姓名
set_cell(c[4], advisor)                         # 指导教师
set_cell(c[5], priority_count)                  # 符合优先条件志愿者总人数 = 7
set_cell(c[6], '1500')                          # 宣传覆盖人数
set_cell(c[7], '是（推荐为优秀团队）')
set_cell(c[8], '2026年暑期依托菏泽单县、曹县、菏泽城区三场高考招生咨询会开展资助政策宣讲；'
              '自主开发AI小程序延伸线上服务；后续进入当地资助中心参与勤工助学实践，'
              '形成“宣讲+数字化+一线实践”的长效服务模式。')

# 合计行：优先条件志愿者明细（col5）
for row in t.rows:
    if row.cells[0].text.strip() == '合计':
        set_cell(row.cells[5],
                 '往届国家或省政府奖学金、励志奖学金，国家助学金获得者(7)人；'
                 '有青年志愿服务经历，参与过山东省青年志愿服务项目大赛的(0)人；'
                 '具备新媒体运营、摄影摄像、文案写作等特长者(4)人')
        break

out2 = os.path.join(BASE, '附件2_青年志愿服务团队汇总表_已填.docx')
doc2.save(out2)
print('保存', out2)
print('完成')
