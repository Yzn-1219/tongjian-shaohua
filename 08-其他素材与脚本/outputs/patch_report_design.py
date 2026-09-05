from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from copy import deepcopy

src = r"C:\Users\Yanzongning\Desktop\童健韶华资料库\校级报告\以动代屏_童健韶华_社会实践报告_最新版（待改）.docx"
doc = Document(src)
paras = doc.paragraphs

er_data_para = None   # (二) 数据尾巴段
er2_para = None       # 二、调研设计与实施 正文段
for p in paras:
    if p.text.startswith("带着这样的设计思路，我们在威海市张村镇黄家皂幼儿园"):
        er_data_para = p
    if p.text.startswith("本次调研采用"):
        er2_para = p

assert er_data_para is not None, "未找到(二)数据段"
assert er2_para is not None, "未找到二、调研设计与实施正文"

def rebuild_para_text(para, text):
    src_rPr = None
    if para.runs and para.runs[0]._r.find(qn('w:rPr')) is not None:
        src_rPr = deepcopy(para.runs[0]._r.find(qn('w:rPr')))
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    run = para.add_run(text)
    if src_rPr is not None:
        run._r.append(deepcopy(src_rPr))
    return para

def copy_para_fmt(src, dst):
    src_pPr = src._p.find(qn('w:pPr'))
    if src_pPr is None:
        return
    dst_pPr = dst._p.get_or_add_pPr()
    for child in list(src_pPr):
        dst_pPr.append(deepcopy(child))

def insert_after(ref_para, text, bold):
    new_p = OxmlElement('w:p')
    ref_para._p.addnext(new_p)
    new_para = Paragraph(new_p, ref_para._parent)
    copy_para_fmt(ref_para, new_para)
    # 继承正文 run 字体（仿宋等）
    src_rPr = None
    if ref_para.runs and ref_para.runs[0]._r.find(qn('w:rPr')) is not None:
        src_rPr = deepcopy(ref_para.runs[0]._r.find(qn('w:rPr')))
    run = new_para.add_run(text)
    if src_rPr is not None:
        run._r.append(deepcopy(src_rPr))
    if bold:
        run.font.bold = True
    return new_para

# ===== 1) (二) 数据尾巴段 → 过渡句 =====
bridge = "带着这样的设计思路，我们走进威海市张村镇三个社区开展实地调研，问卷的题项设计与执行过程详见下节。"
rebuild_para_text(er_data_para, bridge)

# ===== 2) 二、调研设计与实施：导语 + 三小节 =====
intro = '本次调研采用“问卷+访谈”混合设计，力求定量数据与定性观察相互补充。'
rebuild_para_text(er2_para, intro)

subsections = [
    ("（一）问卷设计",
     "家长卷设 15 题，覆盖运动时长、屏幕使用、家长观念与行为、安全顾虑四大维度，并设置逻辑跳转题，减少无效填写；社工卷设 15 题，聚焦社区工作者对辖区儿童体质现状的观察，含环境设施、活动供给与协同难点等维度。两份问卷均经小范围预填与修订，确保表述贴合家长与基层工作者的语言习惯。"),
    ("（二）实地实施",
     "调研于 2026 年 7 月“三下乡”前期集中展开。服务队依托威海市张村镇三个条件各异的社区点位推进：在黄家皂幼儿园，借入园接送时段开展家长座谈与问卷发放；在蓝海社区与西海社区，联合社区工作者通过入户走访、院落宣讲同步收集问卷，并对部分家庭开展深度访谈。为扩大覆盖面，另通过问卷星线上推送，方便不便到场的家长远程填写。执行中，队员逐一说明填答须知、现场答疑，对低龄儿童家庭由家长代填并核对，确保信息真实有效。"),
    ("（三）样本与质控",
     "共回收家长卷 200 份、社工卷 25 份。为保证分析质量，严格按调查原则筛查无效问卷：关键题漏答、前后逻辑矛盾（如自填“几乎不监督”却在监督困难项填“经常”）、规律作答（整卷同一选项）均予剔除。最终家长卷有效 183 份（有效率 91.5%）、社工卷有效 24 份（有效率 96.0%），后续分析均基于有效样本。"),
]

cur = er2_para
for heading, body in subsections:
    hp = insert_after(cur, heading, bold=True)
    bp = insert_after(hp, body, bold=False)
    cur = bp

doc.save(src)
print("DONE: 两处已更新并保存")
